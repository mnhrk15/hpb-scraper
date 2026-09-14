import os
import re
from unittest.mock import patch, MagicMock

import pandas as pd

from app.main.services import check_service
from app.main.services.scraping_service import ScrapingService

# 「髪質改善」のURLエンコード結果
KAMI = '%E9%AB%AA%E8%B3%AA%E6%94%B9%E5%96%84'


class TestBuildFreewordUrl:
    def test_appends_freeword(self, app_context):
        """エリアURLにfreewordクエリを付与し、日本語が%XXエンコードされる。"""
        service = ScrapingService()
        url = service._build_freeword_url(
            'https://beauty.hotpepper.jp/svcSA/macJR/salon/', '髪質改善'
        )
        assert url == f'https://beauty.hotpepper.jp/svcSA/macJR/salon/?freeword={KAMI}'

    def test_none_or_empty_unchanged(self, app_context):
        """freewordがNone/空文字ならbase_urlをそのまま返す（後方互換）。"""
        service = ScrapingService()
        base = 'https://beauty.hotpepper.jp/svcSA/macJR/salon/'
        assert service._build_freeword_url(base, None) == base
        assert service._build_freeword_url(base, '') == base

    def test_preserves_existing_query(self, app_context):
        """既存クエリを保持しつつfreewordを追加する。"""
        service = ScrapingService()
        url = service._build_freeword_url(
            'https://beauty.hotpepper.jp/svcSA/macJR/salon/?searchGender=ALL', '髪質改善'
        )
        assert 'searchGender=ALL' in url
        assert f'freeword={KAMI}' in url


class TestPageUrlConstruction:
    def test_freeword_pagination_url(self, app_context):
        """freeword指定時、page2以降が .../salon/PN{N}.html?freeword=... の正しい形になる。"""
        service = ScrapingService()
        area_url = f'https://beauty.hotpepper.jp/svcSA/macJR/salon/?freeword={KAMI}'
        with patch.object(service, '_get_salon_urls_from_page', return_value=set()) as mock_page:
            list(service._get_all_salon_urls(area_url, 3, 'job123', '髪質改善'))

        called = sorted(c.args[0] for c in mock_page.call_args_list)
        assert called == [
            f'https://beauty.hotpepper.jp/svcSA/macJR/salon/?freeword={KAMI}',
            f'https://beauty.hotpepper.jp/svcSA/macJR/salon/PN2.html?freeword={KAMI}',
            f'https://beauty.hotpepper.jp/svcSA/macJR/salon/PN3.html?freeword={KAMI}',
        ]
        # 壊れたURL（クエリのあとにパスが続く形）が生成されていないこと
        for url in called:
            assert f'{KAMI}/PN' not in url

    def test_no_freeword_backward_compat(self, app_context):
        """freeword未指定時は従来通りの {base}/PN{N}.html を生成する。"""
        service = ScrapingService()
        area_url = 'https://beauty.hotpepper.jp/svcSG/macGE/salon/'
        with patch.object(service, '_get_salon_urls_from_page', return_value=set()) as mock_page:
            list(service._get_all_salon_urls(area_url, 2, 'job123', None))

        called = sorted(c.args[0] for c in mock_page.call_args_list)
        assert called == [
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/',
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/PN2.html',
        ]


class TestGetTotalPagesRedirect:
    def _response(self, url, text='<html><body>no pagination</body></html>'):
        fake = MagicMock()
        fake.url = url
        fake.text = text
        return fake

    def test_freeword_reapplied_when_redirect_drops_query(self, app_context):
        """リダイレクトでfinal_urlからfreewordが落ちても再付与される。"""
        service = ScrapingService()
        # response.url がクエリ無し（=リダイレクトでfreewordが消えた状態）
        fake = self._response('https://beauty.hotpepper.jp/svcSA/macJR/salon/')
        with patch.object(service, '_make_request', return_value=fake):
            total, final = service._get_total_pages(
                f'https://beauty.hotpepper.jp/svcSA/macJR/salon/?freeword={KAMI}',
                'job', '髪質改善',
            )
        assert total == 1  # pagination要素が無いので1ページ
        assert f'freeword={KAMI}' in final

    def test_freeword_preserves_other_query_params(self, app_context):
        """final_urlの他クエリ(searchGender等)を保持しつつfreewordを補う。"""
        service = ScrapingService()
        fake = self._response(
            'https://beauty.hotpepper.jp/svcSA/macJR/salon/?searchGender=ALL'
        )
        with patch.object(service, '_make_request', return_value=fake):
            _, final = service._get_total_pages(
                f'https://beauty.hotpepper.jp/svcSA/macJR/salon/?freeword={KAMI}',
                'job', '髪質改善',
            )
        assert 'searchGender=ALL' in final
        assert f'freeword={KAMI}' in final

    def test_no_freeword_final_url_unchanged(self, app_context):
        """freeword未指定時はfinal_urlを変更しない（後方互換）。"""
        service = ScrapingService()
        url = 'https://beauty.hotpepper.jp/svcSG/macGE/salon/'
        fake = self._response(url)
        with patch.object(service, '_make_request', return_value=fake):
            _, final = service._get_total_pages(url, 'job', None)
        assert final == url


class TestExcelFilename:
    def _df(self):
        return pd.DataFrame({'サロン名': ['サロンA']})

    def test_target_filename_with_freeword(self, app_context):
        """freeword指定時、ファイル名に {area}_{freeword}_ が含まれる。"""
        service = ScrapingService()
        name = service._create_target_excel_file(self._df(), '青山・表参道・原宿', '髪質改善')
        assert name.startswith('青山・表参道・原宿_髪質改善_')
        assert name.endswith('.xlsx')

    def test_excluded_filename_with_freeword(self, app_context):
        """除外リストも 除外リスト_{area}_{freeword}_ の形になる。"""
        service = ScrapingService()
        df = pd.DataFrame({'サロン名': ['A'], 'exclusion_reason': ['EPRP']})
        name = service._create_excluded_excel_file(df, 'エリア', '髪質改善')
        assert name.startswith('除外リスト_エリア_髪質改善_')

    def test_filename_sanitizes_freeword(self, app_context):
        """freeword中のファイル名禁止文字が除去される。"""
        service = ScrapingService()
        name = service._create_target_excel_file(self._df(), 'エリア', '髪/質:改善')
        assert '/' not in name and ':' not in name
        assert name.startswith('エリア_髪質改善_')

    def test_symbol_only_freeword_falls_back(self, app_context):
        """サニタイズ後に空になるfreewordはfreewordなしのファイル名にフォールバックする。"""
        service = ScrapingService()
        name = service._create_target_excel_file(self._df(), 'エリア', '///')
        assert re.match(r'^エリア_\d{8}_\d{6}\.xlsx$', name)

    def test_no_freeword_filename_unchanged(self, app_context):
        """freeword未指定時は従来通り {area}_{timestamp}.xlsx（後方互換）。"""
        service = ScrapingService()
        name = service._create_target_excel_file(self._df(), 'エリア', None)
        assert re.match(r'^エリア_\d{8}_\d{6}\.xlsx$', name)


class TestSalonUrlOrdering:
    """
    電話番号だけ同じ行の重複除去で「どちらを残すか」を再現できるよう、
    サロンURLは一覧の掲載順で返す必要がある。
    """

    def test_urls_are_returned_in_page_order(self, app_context):
        service = ScrapingService()
        pages = {
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/': ['u1', 'u2'],
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/PN2.html': ['u3', 'u4'],
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/PN3.html': ['u5'],
        }
        # _get_all_salon_urls はSSEをyieldしつつURLをreturnするので、戻り値を取り出す
        def consume():
            return (yield from service._get_all_salon_urls(
                'https://beauty.hotpepper.jp/svcSG/macGE/salon/', 3, 'job', None
            ))

        with patch.object(service, '_get_salon_urls_from_page', side_effect=lambda url, job: pages[url]):
            gen = consume()
            try:
                while True:
                    next(gen)
            except StopIteration as stop:
                urls = stop.value

        assert urls == ['u1', 'u2', 'u3', 'u4', 'u5']

    def test_duplicate_urls_across_pages_are_removed_keeping_first(self, app_context):
        service = ScrapingService()
        pages = {
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/': ['u1', 'u2'],
            'https://beauty.hotpepper.jp/svcSG/macGE/salon/PN2.html': ['u2', 'u3'],
        }

        def consume():
            return (yield from service._get_all_salon_urls(
                'https://beauty.hotpepper.jp/svcSG/macGE/salon/', 2, 'job', None
            ))

        with patch.object(service, '_get_salon_urls_from_page', side_effect=lambda url, job: pages[url]):
            gen = consume()
            try:
                while True:
                    next(gen)
            except StopIteration as stop:
                urls = stop.value

        assert urls == ['u1', 'u2', 'u3']


def _row(name, tel, url, excluded=False, reason=''):
    return {
        'サロン名': name, '電話番号': tel, '住所': '住所', 'スタッフ数': 'スタイリスト3人',
        '関連リンク': '', '関連リンク数': 0, 'サロンURL': url,
        'is_excluded': excluded, 'exclusion_reason': reason,
    }


def _consume(generator):
    """SSE文字列を集めつつジェネレータを最後まで回す。"""
    return list(generator)


class TestRunChecks:
    def test_tel_duplicate_keeps_first_row(self, app_context):
        service = ScrapingService()
        df = pd.DataFrame([
            _row('A', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/'),
            _row('B', '0311111111', 'https://beauty.hotpepper.jp/slnH000000002/'),
            _row('C', '03-2222-2222', 'https://beauty.hotpepper.jp/slnH000000003/'),
        ])
        sections = check_service.new_memo_sections()
        app_context.config['STYLIST_CHECK_ENABLED'] = False

        _consume(service._run_checks(df, sections, 'job', None))

        assert df.at[0, 'is_excluded'] is False or not df.at[0, 'is_excluded']
        assert df.at[1, 'is_excluded']
        assert df.at[1, 'exclusion_reason'] == '重複店舗'
        # memoにはグループの全行が「残」「除外」付きで載る
        duplicate_rows = sections[check_service.MEMO_SECTION_DUPLICATE]
        assert [r['サロン名'] for r in duplicate_rows] == ['A', 'B']
        assert [r['備考'] for r in duplicate_rows] == ['残', '除外']

    def test_ng_list_skipped_message(self, app_context):
        service = ScrapingService()
        df = pd.DataFrame([_row('A', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/')])
        app_context.config['STYLIST_CHECK_ENABLED'] = False

        events = _consume(service._run_checks(df, check_service.new_memo_sections(), 'job', None))

        assert any('NGリスト未指定' in e for e in events)

    def test_ng_match_excludes_row(self, app_context):
        service = ScrapingService()
        df = pd.DataFrame([
            _row('A', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/'),
            _row('B', '03-2222-2222', 'https://beauty.hotpepper.jp/slnH000000002/'),
        ])
        sections = check_service.new_memo_sections()
        app_context.config['STYLIST_CHECK_ENABLED'] = False
        ng_index = {'tels': {'0311111111'}, 'ids': set()}

        _consume(service._run_checks(df, sections, 'job', ng_index))

        assert df.at[0, 'is_excluded']
        assert df.at[0, 'exclusion_reason'] == 'NGリスト'
        assert sections[check_service.MEMO_SECTION_NG][0]['備考'] == '一致キー: 電話番号'
        assert not df.at[1, 'is_excluded']

    def test_stylist_check_excludes_and_keeps(self, app_context):
        service = ScrapingService()
        df = pd.DataFrame([
            _row('ひとり店', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/'),
            _row('ふたり店', '03-2222-2222', 'https://beauty.hotpepper.jp/slnH000000002/'),
            _row('取得失敗店', '03-3333-3333', 'https://beauty.hotpepper.jp/slnH000000003/'),
        ])
        sections = check_service.new_memo_sections()
        app_context.config['STYLIST_CHECK_ENABLED'] = True

        def fake_fetch(salon_url, salon_name, job_id):
            if 'H000000001' in salon_url:
                return {'slots': 2, 'assistants': 0, 'placeholders': 1, 'real_count': 1,
                        'placeholder_names': ['店名の枠'], 'names': ['ヤマダ'], 'has_section': True}
            if 'H000000002' in salon_url:
                return {'slots': 2, 'assistants': 0, 'placeholders': 0, 'real_count': 2,
                        'placeholder_names': [], 'names': ['ヤマダ', 'スズキ'], 'has_section': True}
            return None

        with patch.object(service, '_fetch_stylist_info', side_effect=fake_fetch):
            _consume(service._run_checks(df, sections, 'job', None))

        assert df.at[0, 'is_excluded']
        assert df.at[0, 'exclusion_reason'] == 'スタイリスト人数が1名'
        assert not df.at[1, 'is_excluded']
        # 取得失敗は除外せず「要確認」へ
        assert not df.at[2, 'is_excluded']
        review_notes = [r['備考'] for r in sections[check_service.MEMO_SECTION_REVIEW]]
        assert 'スタイリストページ取得失敗' in review_notes

        stylist_rows = sections[check_service.MEMO_SECTION_STYLIST]
        assert stylist_rows[0]['備考'] == '枠2・アシスタント0・人でない枠1（店名の枠）'

    def test_zero_slots_is_not_excluded(self, app_context):
        """見出しは一致するが枠を1つも読めない（セレクタ陳腐化）場合は除外せず要確認へ。"""
        service = ScrapingService()
        df = pd.DataFrame([_row('枠が読めない店', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/')])
        sections = check_service.new_memo_sections()
        app_context.config['STYLIST_CHECK_ENABLED'] = True

        parsed = {'slots': 0, 'assistants': 0, 'placeholders': 0, 'real_count': 0,
                  'placeholder_names': [], 'names': [], 'has_section': True}
        with patch.object(service, '_fetch_stylist_info', return_value=parsed):
            _consume(service._run_checks(df, sections, 'job', None))

        assert not df.at[0, 'is_excluded']
        assert sections[check_service.MEMO_SECTION_STYLIST] == []
        assert sections[check_service.MEMO_SECTION_REVIEW][0]['備考'] == 'スタイリスト枠を読み取れず'

    def test_ng_load_failure_does_not_claim_unspecified(self, app_context):
        """NGリスト指定ありで読み込みに失敗した場合、「未指定」とは言わない。"""
        service = ScrapingService()
        df = pd.DataFrame([_row('A', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/')])
        app_context.config['STYLIST_CHECK_ENABLED'] = False

        events = _consume(
            service._run_checks(df, check_service.new_memo_sections(), 'job', None, ng_requested=True)
        )

        assert not any('NGリスト未指定' in e for e in events)

    def test_keyword_rows_stay_in_target_list(self, app_context):
        service = ScrapingService()
        df = pd.DataFrame([_row('白髪染め専門店 かみ染', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/')])
        sections = check_service.new_memo_sections()
        app_context.config['STYLIST_CHECK_ENABLED'] = False
        app_context.config['CHECK_FLAG_KEYWORDS'] = '白髪染め,かみ染'

        _consume(service._run_checks(df, sections, 'job', None))

        # 要確認は除外しない（人の判断が要るものを機械が黙って落とさない）
        assert not df.at[0, 'is_excluded']
        assert sections[check_service.MEMO_SECTION_REVIEW][0]['備考'] == '店名キーワード: 白髪染め／かみ染'

    def test_existing_exclusions_go_to_memo_sections(self, app_context):
        service = ScrapingService()
        df = pd.DataFrame([
            _row('EPRP店', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/', True, 'EPRP'),
            _row('リンク多い店', '03-2222-2222', 'https://beauty.hotpepper.jp/slnH000000002/', True, '関連リンク数'),
            _row('1人店', '03-3333-3333', 'https://beauty.hotpepper.jp/slnH000000003/', True, 'スタイリスト人数が1名'),
        ])
        sections = check_service.new_memo_sections()
        app_context.config['STYLIST_CHECK_ENABLED'] = False

        _consume(service._run_checks(df, sections, 'job', None))

        assert [r['サロン名'] for r in sections[check_service.MEMO_SECTION_EPRP]] == ['EPRP店']
        assert [r['サロン名'] for r in sections[check_service.MEMO_SECTION_RELATED_LINKS]] == ['リンク多い店']
        assert [r['サロン名'] for r in sections[check_service.MEMO_SECTION_STYLIST]] == ['1人店']

    def test_excluded_rows_are_not_fetched_for_stylists(self, app_context):
        """既存の除外で落ちた行にはスタイリストタブを取りに行かない。"""
        service = ScrapingService()
        df = pd.DataFrame([
            _row('除外済み', '03-1111-1111', 'https://beauty.hotpepper.jp/slnH000000001/', True, 'EPRP'),
            _row('対象', '03-2222-2222', 'https://beauty.hotpepper.jp/slnH000000002/'),
        ])
        app_context.config['STYLIST_CHECK_ENABLED'] = True

        with patch.object(service, '_fetch_stylist_info', return_value=None) as mock_fetch:
            _consume(service._run_checks(df, check_service.new_memo_sections(), 'job', None))

        fetched_urls = [call.args[0] for call in mock_fetch.call_args_list]
        assert fetched_urls == ['https://beauty.hotpepper.jp/slnH000000002/']


class TestMemoExcelFile:
    def test_filename_and_sections(self, app_context):
        import openpyxl

        service = ScrapingService()
        sections = check_service.new_memo_sections()
        sections[check_service.MEMO_SECTION_STYLIST].append({
            'サロン名': 'A', '電話番号': '03-1111-1111', '住所': '住所', 'スタッフ数': 'スタイリスト2人',
            '関連リンク': '', '関連リンク数': 0, 'サロンURL': 'https://beauty.hotpepper.jp/slnH000000001/',
            '備考': '枠2・アシスタント0・人でない枠1（店名）',
        })

        file_name = service._create_memo_excel_file(sections, '秋田', None)
        assert re.match(r'^対応メモ_秋田_\d{8}_\d{6}\.xlsx$', file_name)

        path = os.path.join(app_context.config['OUTPUT_DIR'], file_name)
        sheet = openpyxl.load_workbook(path).active
        assert sheet.title == '秋田'

        rows = [[c.value for c in row] for row in sheet.iter_rows()]
        headings = [r[0] for r in rows if r[0] and str(r[0]).startswith('■')]
        assert headings == check_service.MEMO_SECTION_ORDER

        # 該当行は7列＋備考（空文字はExcel上は空セル＝読み戻すとNone）
        assert rows[1][:8] == [
            'A', '03-1111-1111', '住所', 'スタイリスト2人', None, 0,
            'https://beauty.hotpepper.jp/slnH000000001/', '枠2・アシスタント0・人でない枠1（店名）',
        ]
        # 該当なしの区分には「該当なし」が入る
        assert '該当なし' in [r[0] for r in rows]

    def test_filename_with_freeword(self, app_context):
        service = ScrapingService()
        file_name = service._create_memo_excel_file(
            check_service.new_memo_sections(), 'エリア', '髪質改善'
        )
        assert file_name.startswith('対応メモ_エリア_髪質改善_')

    def test_long_area_name_is_truncated_for_sheet(self, app_context):
        import openpyxl

        service = ScrapingService()
        area_name = 'あ' * 40
        file_name = service._create_memo_excel_file(check_service.new_memo_sections(), area_name, None)
        sheet = openpyxl.load_workbook(os.path.join(app_context.config['OUTPUT_DIR'], file_name)).active
        assert sheet.title == 'あ' * 31
