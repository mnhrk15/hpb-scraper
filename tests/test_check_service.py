import os

import openpyxl
import pytest

from app.main.services import check_service as check

FIXTURE_DIR = os.path.join(os.path.dirname(__file__), 'fixtures')


def read_fixture(name):
    with open(os.path.join(FIXTURE_DIR, name), encoding='utf-8') as f:
        return f.read()


class TestParseStylistPage:
    """
    2026-09-14に実サイトから取得した3店舗のスタイリストタブで固定する。
    3店舗とも2026-07の人手は「スタイリスト人数が1名」で除外している。
    """

    def test_loews_counts_one_real_stylist(self):
        """実在の1人＋サロン名の枠1つ → 実1人（サロン名の枠は数えない）。"""
        parsed = check.parse_stylist_page(
            read_fixture('stylist_loews.html'), 'LOEWS SALON 　新小岩　【ロウズ】'
        )
        assert parsed['slots'] == 2
        assert parsed['placeholders'] == 1
        assert parsed['placeholder_names'] == ['LOEWS SALON']
        assert parsed['names'] == ['ナカムラ ユウスケ']
        assert parsed['real_count'] == 1
        assert parsed['has_section'] is True

    def test_kamisome_counts_zero_real_stylists(self):
        """2枠とも店名（「かみ染 小岩」「かみ染１　小岩」）→ 実0人。"""
        parsed = check.parse_stylist_page(
            read_fixture('stylist_kamisome.html'), '白髪染め専門店　かみ染　小岩'
        )
        assert parsed['slots'] == 2
        assert parsed['placeholders'] == 2
        assert parsed['real_count'] == 0

    def test_roi_counts_zero_real_stylists(self):
        """1枠が店名の略称（roi）→ 実0人。サロン名は()や【】の中も候補にする。"""
        parsed = check.parse_stylist_page(
            read_fixture('stylist_roi.html'), 'HAIR COLOR SALON roi【ヘアカラーサロン　ロイ】'
        )
        assert parsed['slots'] == 1
        assert parsed['placeholders'] == 1
        assert parsed['real_count'] == 0

    def test_reserve_button_blocks_are_not_counted(self):
        """div.w166.mHA には「指名して予約する」等の枠も混ざるので枠数に数えない。"""
        parsed = check.parse_stylist_page(
            read_fixture('stylist_loews.html'), 'LOEWS SALON 　新小岩　【ロウズ】'
        )
        # 元のHTMLでは div.w166.mHA は6個あるが、名前のある枠は2個
        assert parsed['slots'] == 2

    def test_assistant_is_not_counted(self):
        """役職にアシスタントを含む枠はスタイリストとして数えない。"""
        html = (
            '<html><body><h1>テストサロンのスタイリスト</h1>'
            '<div class="w166 mHA"><p class="mT10 fs16 b">ヤマダ タロウ</p>'
            '<div class="mT5 fs10"><span class="fgPink">アシスタント</span></div></div>'
            '<div class="w166 mHA"><p class="mT10 fs16 b">スズキ ハナコ</p>'
            '<div class="mT5 fs10"><span class="fgPink">スタイリスト</span></div></div>'
            '</body></html>'
        )
        parsed = check.parse_stylist_page(html, 'テストサロン')
        assert parsed['slots'] == 2
        assert parsed['assistants'] == 1
        assert parsed['real_count'] == 1

    def test_placeholder_words(self):
        """「ゲスト」「BARBER stylist」などの定型語の枠は人として数えない。"""
        html = (
            '<html><body><h1>テストサロンのスタイリスト</h1>'
            '<div class="w166 mHA"><p class="mT10 fs16 b">ゲスト</p></div>'
            '<div class="w166 mHA"><p class="mT10 fs16 b">BARBER stylist</p></div>'
            '<div class="w166 mHA"><p class="mT10 fs16 b">ヘアの ご予約4</p></div>'
            '<div class="w166 mHA"><p class="mT10 fs16 b">タナカ ケンジ</p></div>'
            '</body></html>'
        )
        parsed = check.parse_stylist_page(html, 'テストサロン')
        assert parsed['placeholders'] == 3
        assert parsed['real_count'] == 1

    def test_no_stylist_section(self):
        """スタイリストの枠も見出しも無ければ has_section が False（＝除外しない）。"""
        parsed = check.parse_stylist_page('<html><body><h1>サロン紹介</h1></body></html>', 'テストサロン')
        assert parsed['has_section'] is False
        assert parsed['real_count'] == 0


class TestStylistPagination:
    def test_single_page(self):
        assert check.parse_stylist_total_pages(read_fixture('stylist_loews.html')) == 1

    def test_multi_page(self):
        assert check.parse_stylist_total_pages(read_fixture('stylist_multipage.html')) == 3

    def test_missing_pagination_defaults_to_one(self):
        assert check.parse_stylist_total_pages('<html><body></body></html>') == 1


class TestBuildStylistUrl:
    def test_first_page(self):
        assert check.build_stylist_url('https://beauty.hotpepper.jp/slnH000064475/') == \
            'https://beauty.hotpepper.jp/slnH000064475/stylist/'

    def test_without_trailing_slash(self):
        assert check.build_stylist_url('https://beauty.hotpepper.jp/slnH000064475') == \
            'https://beauty.hotpepper.jp/slnH000064475/stylist/'

    def test_second_page(self):
        """2ページ目以降は /stylist/PN{n}.html（2026-09-14に実サイトで確認）。"""
        assert check.build_stylist_url('https://beauty.hotpepper.jp/slnH000064475/', 2) == \
            'https://beauty.hotpepper.jp/slnH000064475/stylist/PN2.html'

    def test_empty_url(self):
        assert check.build_stylist_url('') == ''


class TestNormalizeTel:
    @pytest.mark.parametrize('value,expected', [
        ('03-1234-5678', '0312345678'),
        ('０３－１２３４－５６７８', '0312345678'),  # 全角はNFKCで半角に寄せて拾う
        ('(03) 1234 5678', '0312345678'),
        ('', None),
        (None, None),
        ('電話なし', None),
    ])
    def test_normalize(self, value, expected):
        assert check.normalize_tel(value) == expected


class TestExtractSalonId:
    def test_extracts_id(self):
        assert check.extract_salon_id('https://beauty.hotpepper.jp/slnH000681771/') == 'H000681771'

    def test_short_id(self):
        assert check.extract_salon_id('https://beauty.hotpepper.jp/slnH00039042') == 'H00039042'

    def test_no_match(self):
        assert check.extract_salon_id('https://example.com/') is None


class TestDedupeByTel:
    def test_keeps_first_of_group(self):
        """同じ電話番号のグループは先頭（一覧の出現順）を残す。"""
        dropped, groups = check.dedupe_by_tel([
            (0, '03-1111-1111'),
            (1, '0311111111'),
            (2, '03-2222-2222'),
            (3, '03-1111-1111'),
        ])
        assert dropped == {1, 3}
        assert len(groups) == 1
        tel, members = groups[0]
        assert tel == '0311111111'
        # memoにはグループの全行が載り、残した行が分かる
        assert members == [(0, True), (1, False), (3, False)]

    def test_empty_tel_is_ignored(self):
        dropped, groups = check.dedupe_by_tel([(0, ''), (1, None), (2, '電話なし')])
        assert dropped == set()
        assert groups == []

    def test_no_duplicates(self):
        dropped, groups = check.dedupe_by_tel([(0, '03-1111-1111'), (1, '03-2222-2222')])
        assert dropped == set()
        assert groups == []


class TestFlagKeywords:
    def test_default_keywords(self):
        keywords = check.parse_flag_keywords('')
        assert 'カラー専門' in keywords

    def test_parse_comma_separated(self):
        assert check.parse_flag_keywords('A, B ,, C') == ['A', 'B', 'C']

    def test_case_and_width_insensitive(self):
        """大文字小文字・全半角を無視して店名に含むか判定する。"""
        keywords = check.parse_flag_keywords('NEW OPEN,カラー専門')
        assert check.find_flag_keywords('ＮＥＷ　ｏｐｅｎ　サロン', keywords) == ['NEW OPEN']
        assert check.find_flag_keywords('カラー専門店ABC', keywords) == ['カラー専門']

    def test_no_match(self):
        keywords = check.parse_flag_keywords('NEW OPEN')
        assert check.find_flag_keywords('ヘアサロン ABC', keywords) == []


class TestBuildNgIndex:
    def _make_workbook(self, path):
        """見出し行の位置も列の並びも違う2シートのNGリストを作る。"""
        workbook = openpyxl.Workbook()

        sheet1 = workbook.active
        sheet1.title = '東京都'
        sheet1.append(['店名', '住所', '電話番号', 'HPBリンク', '備考欄'])
        sheet1.append(['サロンA', '東京都…', '03-1234-5678', 'https://beauty.hotpepper.jp/slnH000111111/', None])
        sheet1.append(['サロンB', '東京都…', '090-1111-2222\n03-9999-8888', None, None])

        sheet2 = workbook.create_sheet('北海道')
        sheet2.append([None, None, None, None])          # 空の1行目（見出しの位置が違う）
        sheet2.append(['店名', '住所', '電話番号', 'URL'])
        sheet2.append(['サロンC', '北海道…', None, 'https://beauty.hotpepper.jp/slnH000222222/'])
        sheet2.append(['サロンD', '北海道…', '011-222-3333 / 011-222-4444', '掲載なし'])

        workbook.save(path)

    def test_collects_tels_and_ids_across_sheets(self, tmp_path):
        path = str(tmp_path / 'ng.xlsx')
        self._make_workbook(path)
        index = check.build_ng_index(path)

        assert 'H000111111' in index['ids']
        assert 'H000222222' in index['ids']
        assert '0312345678' in index['tels']
        # セル内に改行や「/」で複数入っていても分けて拾う
        assert '09011112222' in index['tels']
        assert '0399998888' in index['tels']
        assert '0112223333' in index['tels']
        assert '0112224444' in index['tels']

    def test_match_ng_by_tel_and_id(self, tmp_path):
        path = str(tmp_path / 'ng.xlsx')
        self._make_workbook(path)
        index = check.build_ng_index(path)

        assert check.match_ng('03-1234-5678', 'https://beauty.hotpepper.jp/slnH000999999/', index) == '電話番号'
        assert check.match_ng('0000000000', 'https://beauty.hotpepper.jp/slnH000222222/', index) == 'サロンID'
        assert check.match_ng('0000000000', 'https://beauty.hotpepper.jp/slnH000999999/', index) is None

    def test_match_ng_without_index(self):
        assert check.match_ng('03-1234-5678', 'https://beauty.hotpepper.jp/slnH000111111/', None) is None


class TestSafeSheetName:
    def test_removes_forbidden_characters(self):
        assert check.safe_sheet_name('a[b]c:d*e?f/g\\h') == 'abcdefgh'

    def test_truncates_to_31_chars(self):
        assert len(check.safe_sheet_name('あ' * 40)) == 31

    def test_empty_falls_back(self):
        assert check.safe_sheet_name('') == 'シート1'


class TestMemoSections:
    def test_order_matches_delivered_memo(self):
        """区分の並びは2026-07の人手の納品物（対応メモ_20260723.xlsx）と同じ。"""
        assert check.MEMO_SECTION_ORDER[:6] == [
            '■スタイリスト人数が1名',
            '■関連リンク4店舗以上',
            '■NGリストに記載のある店舗',
            '■EPRP店舗、掲載終了店舗',
            '■重複店舗 ：電話番号・HPBリンク',
            '■その他',
        ]
        # 末尾の「要確認」だけが今回の追加
        assert check.MEMO_SECTION_ORDER[6].startswith('■要確認')

    def test_new_memo_sections_is_empty(self):
        sections = check.new_memo_sections()
        assert list(sections.keys()) == check.MEMO_SECTION_ORDER
        assert all(rows == [] for rows in sections.values())


class TestSalonNameInStylistName:
    """
    店名を名前に付けて登録するサロンがある（`天野朝飛 few.新小岩`）。
    店名を含むだけで「人でない枠」にすると、実在のスタイリストを落としてしまう。
    """

    def _page(self, names):
        slots = ''.join(
            f'<div class="w166 mHA"><p class="mT10 fs16 b">{name}</p></div>' for name in names
        )
        return f'<html><body><h1>few. 新小岩 【フュー】のスタイリスト</h1>{slots}</body></html>'

    def test_personal_name_plus_salon_name_counts_as_person(self):
        parsed = check.parse_stylist_page(
            self._page(['天野朝飛 few.新小岩', 'Ameri few.新小岩UP', '寺木拓海 few.新小岩']),
            'few. 新小岩 【フュー】',
        )
        assert parsed['placeholders'] == 0
        assert parsed['real_count'] == 3

    def test_salon_name_alone_is_still_a_placeholder(self):
        """店名だけ、または店名＋枝番（UP）の枠は人として数えない。"""
        parsed = check.parse_stylist_page(
            self._page(['few. 新小岩', 'few.新小岩UP']), 'few. 新小岩 【フュー】'
        )
        assert parsed['placeholders'] == 2
        assert parsed['real_count'] == 0

    def test_short_personal_name_still_counts(self):
        """3文字の名前（mai）は人として数える。"""
        parsed = check.parse_stylist_page(
            self._page(['mai few.新小岩']), 'few. 新小岩 【フュー】'
        )
        assert parsed['real_count'] == 1
