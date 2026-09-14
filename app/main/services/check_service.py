"""
打電リストチェックの判定ロジック。

HTTP取得とExcel出力から独立した純関数だけを置く（テスト対象）。
ScrapingServiceはここの関数を呼び出し、SSEの文言とファイル出力だけを担当する。
"""

import re
import unicodedata
from collections import OrderedDict

from bs4 import BeautifulSoup

# 営業対象リスト・除外リスト・対応メモで共通の7列（後工程が読む形。変更しない）
LIST_COLUMNS = ['サロン名', '電話番号', '住所', 'スタッフ数', '関連リンク', '関連リンク数', 'サロンURL']

# 対応メモの区分。2026-07の人手の納品物（対応メモ_20260723.xlsx）と同じ並び。
# 末尾の「■要確認」だけが今回の追加（機械が黙って落とさないものの受け皿）。
MEMO_SECTION_STYLIST = '■スタイリスト人数が1名'
MEMO_SECTION_RELATED_LINKS = '■関連リンク4店舗以上'
MEMO_SECTION_NG = '■NGリストに記載のある店舗'
MEMO_SECTION_EPRP = '■EPRP店舗、掲載終了店舗'
MEMO_SECTION_DUPLICATE = '■重複店舗 ：電話番号・HPBリンク'
MEMO_SECTION_OTHER = '■その他'
MEMO_SECTION_REVIEW = '■要確認（専門店・シェアサロン・新規オープン・取得失敗）'

MEMO_SECTION_ORDER = [
    MEMO_SECTION_STYLIST,
    MEMO_SECTION_RELATED_LINKS,
    MEMO_SECTION_NG,
    MEMO_SECTION_EPRP,
    MEMO_SECTION_DUPLICATE,
    MEMO_SECTION_OTHER,
    MEMO_SECTION_REVIEW,
]

# 除外理由 → 対応メモの区分
REASON_TO_SECTION = {
    'スタイリスト人数が1名': MEMO_SECTION_STYLIST,
    '関連リンク数': MEMO_SECTION_RELATED_LINKS,
    'NGリスト': MEMO_SECTION_NG,
    'EPRP': MEMO_SECTION_EPRP,
    '重複店舗': MEMO_SECTION_DUPLICATE,
}

# スタイリストタブの既定セレクタ（selectors.jsonの stylist_page で上書きできる）
DEFAULT_STYLIST_SELECTORS = {
    'slot': 'div.w166.mHA',
    'name': 'p.mT10.fs16.b',
    'kana': 'p.db.fs10.fgBbrown',
    'role': 'div.mT5.fs10 span.fgPink',
    'pagination': 'p.pa.bottom0.right0',
    'heading': 'h1',
}

# 枠の名前がこれらを含む場合は「人でない枠」とみなす。
# メニュー名（「ヘアの ご予約4」）・店名・「BARBER stylist」などを拾う。
PLACEHOLDER_WORDS = (
    'スタイリスト', 'stylist', 'ゲスト', 'guest', '予約', '指名なし',
    'フリー', '担当', '店舗', 'サロン', 'salon',
)

# 役職にこれを含む枠はスタイリストとして数えない
ASSISTANT_WORD = 'アシスタント'

# 店名を含む枠名から店名を除いた残りが、人名として成立するとみなす最小文字数。
# 実測: 「天野朝飛」「Ameri」「mai」は人、「UP」は枝番の飾り。
MIN_PERSONAL_NAME_LENGTH = 3

# 店名の要確認キーワード既定値（config.CHECK_FLAG_KEYWORDS で上書き）
DEFAULT_FLAG_KEYWORDS = (
    'カラー専門', '白髪染め', 'かみ染', 'シェアサロン', 'メンズ専門',
    'NEW OPEN', 'OPEN予定', '閉店',
)


def new_memo_sections():
    """対応メモの区分を実物と同じ順で空のまま用意する。"""
    return OrderedDict((section, []) for section in MEMO_SECTION_ORDER)


def normalize_tel(value):
    """
    電話番号から数字以外を除去する。数字が無ければNoneを返す。
    NGリストには全角数字の番号が混ざるので先にNFKCで半角に寄せる。
    """
    text = unicodedata.normalize('NFKC', str(value or ''))
    digits = re.sub(r'\D', '', text)
    return digits or None


def extract_salon_id(url):
    """サロンURLからHPBのサロンID（slnH…のH以降）を取り出す。見つからなければNone。"""
    match = re.search(r'sln(H\d+)', str(url or ''))
    return match.group(1) if match else None


def normalize_name(value):
    """
    店名・スタイリスト名の比較用に正規化する。
    NFKC → 小文字 → 英数字とかな漢字だけ残す（空白・記号・数字を除去）。

    数字を落とすのは「かみ染１　小岩」と「かみ染　小岩」を同一視するため。
    2026-07の人手はこの2枠をどちらも店名の枠とみなしている。
    """
    text = unicodedata.normalize('NFKC', str(value or '')).lower()
    return ''.join(ch for ch in text if ch.isalnum() and not ch.isdigit())


def salon_name_candidates(salon_name):
    """
    サロン名の比較候補を作る。
    正式名称に加えて【】・()・（）の中身と、その手前の部分も候補にする。
    例: 'HAIR COLOR SALON roi【ヘアカラーサロン　ロイ】' → 全体 / ヘアカラーサロンロイ / haircolorsalonroi
    """
    raw = str(salon_name or '')
    candidates = [raw]
    for groups in re.findall(r'【(.+?)】|\((.+?)\)|（(.+?)）', raw):
        candidates.extend(g for g in groups if g)
    candidates.append(re.split(r'[【(（]', raw)[0])

    normalized = []
    for candidate in candidates:
        value = normalize_name(candidate)
        # 1文字の候補は誤爆するので使わない
        if len(value) >= 2 and value not in normalized:
            normalized.append(value)
    return normalized


def is_placeholder_name(name, salon_name_variants):
    """
    枠の名前が「人でない枠」かどうかを判定し、(判定, 理由) を返す。
    理由は対応メモの備考に出す用。
    """
    normalized = normalize_name(name)
    if not normalized:
        return True, '名前なし'

    for word in PLACEHOLDER_WORDS:
        if normalize_name(word) in normalized:
            return True, f'定型語「{word}」'

    if len(normalized) >= 2:
        for variant in salon_name_variants:
            # 名前がサロン名に収まる（「かみ染 小岩」「roi」など）＝店名だけの枠
            if normalized in variant:
                return True, '店名と同一'
            # 名前がサロン名を含む場合は、店名を除いた残りが人名として成立するかで分ける。
            # 「天野朝飛 few.新小岩」のように氏名＋店名で登録する店があり、
            # 含むだけで人でない枠にすると実在のスタイリストを落としてしまう。
            # 残りが2文字以下なら枝番などの飾り（「few.新小岩UP」）とみなす。
            if variant in normalized and len(normalized.replace(variant, '', 1)) < MIN_PERSONAL_NAME_LENGTH:
                return True, '店名と同一'

    return False, ''


def parse_stylist_page(html, salon_name, selectors=None):
    """
    スタイリストタブのHTMLを解析し、実スタイリスト数を数える。

    戻り値のdict:
      slots            枠数（名前のある枠だけ。「指名して予約する」等の枠は含めない）
      assistants       アシスタントの枠数
      placeholders     人でない枠の数
      real_count       実スタイリスト数 = slots - assistants - placeholders
      placeholder_names 人でない枠として除いた名前
      names            実スタイリストとして数えた名前
      has_section      スタイリスト一覧のページとして成立しているか
    """
    sel = dict(DEFAULT_STYLIST_SELECTORS)
    if selectors:
        sel.update({k: v for k, v in selectors.items() if v})

    soup = BeautifulSoup(html or '', 'html.parser')
    variants = salon_name_candidates(salon_name)

    # div.w166.mHA には「指名して予約する」ボタンや空のセルも混ざるため、
    # 名前の要素を持つ枠だけをスタイリストの枠として扱う。
    slots = [s for s in soup.select(sel['slot']) if s.select_one(sel['name'])]

    assistants = 0
    placeholders = 0
    placeholder_names = []
    names = []

    for slot in slots:
        name_el = slot.select_one(sel['name'])
        role_el = slot.select_one(sel['role'])
        name = name_el.get_text(strip=True) if name_el else ''
        role = role_el.get_text(strip=True) if role_el else ''

        if ASSISTANT_WORD in role:
            assistants += 1
            continue

        is_placeholder, _reason = is_placeholder_name(name, variants)
        if is_placeholder:
            placeholders += 1
            placeholder_names.append(name)
        else:
            names.append(name)

    heading_el = soup.select_one(sel['heading'])
    heading_text = heading_el.get_text(strip=True) if heading_el else ''
    has_section = bool(slots) or 'スタイリスト' in heading_text

    return {
        'slots': len(slots),
        'assistants': assistants,
        'placeholders': placeholders,
        'real_count': len(names),
        'placeholder_names': placeholder_names,
        'names': names,
        'has_section': has_section,
    }


def parse_stylist_total_pages(html, selectors=None):
    """スタイリストタブの総ページ数を返す。表記が読めなければ1。"""
    sel = dict(DEFAULT_STYLIST_SELECTORS)
    if selectors:
        sel.update({k: v for k, v in selectors.items() if v})

    soup = BeautifulSoup(html or '', 'html.parser')
    element = soup.select_one(sel['pagination'])
    text = element.get_text(strip=True) if element else soup.get_text(' ', strip=True)

    match = re.search(r'\d+\s*/\s*(\d+)ページ', text)
    return int(match.group(1)) if match else 1


def build_stylist_url(salon_url, page=1):
    """
    サロンURLからスタイリストタブのURLを組み立てる。
    2ページ目以降は {サロンURL}stylist/PN{n}.html（2026-09-14に実サイトで確認）。
    """
    base = str(salon_url or '').split('?')[0].rstrip('/')
    if not base:
        return ''
    if page <= 1:
        return f'{base}/stylist/'
    return f'{base}/stylist/PN{page}.html'


def stylist_memo_note(parsed):
    """スタイリスト判定の備考文を作る。"""
    note = (
        f"枠{parsed['slots']}・アシスタント{parsed['assistants']}"
        f"・人でない枠{parsed['placeholders']}"
    )
    if parsed['placeholder_names']:
        note += '（' + '／'.join(parsed['placeholder_names']) + '）'
    return note


def dedupe_by_tel(rows):
    """
    電話番号だけが同じ行をまとめ、先頭を残して残りを除外対象として返す。

    rows は (index, 電話番号) のイテラブル。indexは一覧の出現順に並んでいる前提。
    戻り値 (dropped, groups):
      dropped 除外する index の集合
      groups  [(電話番号, [(index, 残すか)…])…]（2件以上のグループのみ。対応メモ用に全行を載せる）
    """
    by_tel = OrderedDict()
    for index, tel in rows:
        normalized = normalize_tel(tel)
        if not normalized:
            continue
        by_tel.setdefault(normalized, []).append(index)

    dropped = set()
    groups = []
    for tel, indexes in by_tel.items():
        if len(indexes) < 2:
            continue
        dropped.update(indexes[1:])
        groups.append((tel, [(i, i == indexes[0]) for i in indexes]))
    return dropped, groups


def parse_flag_keywords(value):
    """CHECK_FLAG_KEYWORDS設定（カンマ区切り文字列 or リスト）をリストにする。"""
    if not value:
        return list(DEFAULT_FLAG_KEYWORDS)
    if isinstance(value, str):
        items = [v.strip() for v in value.split(',')]
    else:
        items = [str(v).strip() for v in value]
    return [v for v in items if v]


def find_flag_keywords(salon_name, keywords):
    """店名に含まれる要確認キーワードを返す（大文字小文字・全半角を無視）。"""
    normalized = normalize_name(salon_name)
    if not normalized:
        return []
    return [kw for kw in keywords if normalize_name(kw) and normalize_name(kw) in normalized]


def build_ng_index(xlsx_path_or_file):
    """
    打電NGリストのxlsxから電話番号とサロンIDの索引を作る。

    シートごとに見出し行の位置も列の有無も違うため、全シート・全セルを走査して
    文字列から拾う（列名に依存しない）。
    戻り値 {'tels': set, 'ids': set}
    """
    import openpyxl

    tels = set()
    ids = set()

    workbook = openpyxl.load_workbook(xlsx_path_or_file, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    text = unicodedata.normalize('NFKC', str(cell))

                    for salon_id in re.findall(r'sln(H\d+)', text):
                        ids.add(salon_id)

                    # セル内に改行や「/」で複数の番号が入ることがあるので先に切り出す。
                    # 空白類に改行を含めると2つの番号が1つに繋がってしまうため、
                    # 区切りにならない空白（半角スペース・タブ）だけを許す。
                    for chunk in re.findall(r'[\d\-()（） \t]{9,}', text):
                        normalized = normalize_tel(chunk)
                        # 日付("2026-07-23 10:00" → "2026072310")などを拾わないよう、
                        # 国内の電話番号と同じく先頭が0のものだけを採用する。
                        if (
                            normalized
                            and 10 <= len(normalized) <= 11
                            and normalized.startswith('0')
                        ):
                            tels.add(normalized)
    finally:
        workbook.close()

    return {'tels': tels, 'ids': ids}


def match_ng(tel, salon_url, ng_index):
    """
    NGリストとの一致を判定し、一致キー（'電話番号' / 'サロンID'）を返す。
    一致しなければNone。店名は表記揺れで誤爆するので使わない。
    """
    if not ng_index:
        return None

    normalized = normalize_tel(tel)
    if normalized and normalized in ng_index.get('tels', ()):
        return '電話番号'

    salon_id = extract_salon_id(salon_url)
    if salon_id and salon_id in ng_index.get('ids', ()):
        return 'サロンID'

    return None


def safe_sheet_name(name):
    """Excelのシート名として使える形にする（禁止文字を除き31文字に切る）。"""
    cleaned = re.sub(r'[\[\]:*?/\\]', '', str(name or '')).strip()
    return (cleaned or 'シート1')[:31]
